import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, Alignment
import os

def analyze_retention(file_path):
    """
    分析付费用户的留存行为
    1. 计算1天、7天、30天留存率（按自然周分组）
    2. 分析留存天数频数分布
    """
    
    # 读取Excel文件
    print("正在读取Excel文件...")
    df = pd.read_excel(file_path)
    
    # 显示数据基本信息
    print(f"数据总行数: {len(df)}")
    print(f"列名: {df.columns.tolist()}")
    
    # 转换时间列为datetime类型
    df['注册时间'] = pd.to_datetime(df['注册时间'])
    df['最后登录时间'] = pd.to_datetime(df['最后登录时间'], errors='coerce')
    
    # 只保留有注册时间和最后登录时间的用户
    df = df.dropna(subset=['注册时间', '最后登录时间'])
    print(f"有效付费用户数: {len(df)}")
    
    # 计算注册日期（去掉时间部分）
    df['注册日期'] = df['注册时间'].dt.date
    
    # 计算最后登录日期（去掉时间部分）
    df['最后登录日期'] = df['最后登录时间'].dt.date
    
    # 计算留存天数
    df['留存天数'] = (df['最后登录时间'] - df['注册时间']).dt.days
    
    # 确保留存天数非负
    df = df[df['留存天数'] >= 0]
    print(f"有效留存数据用户数: {len(df)}")
    
    # ========== 表单1：按自然周留存率分析 ==========
    print("\n=== 按自然周留存率分析 ===")
    
    # 计算每个用户的注册周
    df['注册周'] = df['注册时间'].dt.isocalendar().week
    df['注册年份'] = df['注册时间'].dt.year
    
    # 创建周标识（年份-周数）
    df['周标识'] = df['注册年份'].astype(str) + '-W' + df['注册周'].astype(str).str.zfill(2)
    
    # 计算每周的留存率
    weekly_retention_results = []
    
    # 按周分组
    weekly_groups = df.groupby('周标识')
    
    for week_id, week_data in weekly_groups:
        reg_count = len(week_data)
        
        # 计算留存用户数
        d1_retained = len(week_data[week_data['留存天数'] >= 1])
        d7_retained = len(week_data[week_data['留存天数'] >= 7])
        d30_retained = len(week_data[week_data['留存天数'] >= 30])
        
        # 计算留存率
        d1_rate = d1_retained / reg_count if reg_count > 0 else 0
        d7_rate = d7_retained / reg_count if reg_count > 0 else 0
        d30_rate = d30_retained / reg_count if reg_count > 0 else 0
        
        # 计算周的日期区间
        week_start = week_data['注册时间'].min().strftime('%m%d')
        week_end = week_data['注册时间'].max().strftime('%m%d')
        year_short = str(week_data['注册时间'].min().year)[-2:]  # 取年份后两位
        date_range = f"{year_short}-{week_start}-{week_end}"
        
        weekly_retention_results.append({
            '自然周': date_range,
            '注册人数': reg_count,
            'D1留存人数': d1_retained,
            'D7留存人数': d7_retained,
            'D30留存人数': d30_retained,
            'D1留存率': d1_rate,
            'D7留存率': d7_rate,
            'D30留存率': d30_rate
        })
    
    # 转换为DataFrame并按周排序
    weekly_retention_df = pd.DataFrame(weekly_retention_results)
    weekly_retention_df = weekly_retention_df.sort_values('自然周')
    
    # 计算总体平均留存率
    total_registrations = weekly_retention_df['注册人数'].sum()
    total_d1_retained = weekly_retention_df['D1留存人数'].sum()
    total_d7_retained = weekly_retention_df['D7留存人数'].sum()
    total_d30_retained = weekly_retention_df['D30留存人数'].sum()
    
    overall_d1_rate = total_d1_retained / total_registrations if total_registrations > 0 else 0
    overall_d7_rate = total_d7_retained / total_registrations if total_registrations > 0 else 0
    overall_d30_rate = total_d30_retained / total_registrations if total_registrations > 0 else 0
    
    print(f"总注册人数: {total_registrations}")
    print(f"1天留存人数: {total_d1_retained}")
    print(f"7天留存人数: {total_d7_retained}")
    print(f"30天留存人数: {total_d30_retained}")
    print(f"总体1天留存率: {overall_d1_rate:.4f} ({overall_d1_rate*100:.2f}%)")
    print(f"总体7天留存率: {overall_d7_rate:.4f} ({overall_d7_rate*100:.2f}%)")
    print(f"总体30天留存率: {overall_d30_rate:.4f} ({overall_d30_rate*100:.2f}%)")
    
    # ========== 表单2：留存天数频数分布 ==========
    print("\n=== 留存天数频数分布分析 ===")
    
    # 定义留存天数区间
    def categorize_retention_days(days):
        if days == 0:
            return '当天'
        elif days == 1:
            return '次日'
        elif 2 <= days <= 7:
            return '2-7天'
        elif 8 <= days <= 30:
            return '8-30天'
        else:
            return '30天以上'
    
    # 添加留存天数分类
    df['留存天数分类'] = df['留存天数'].apply(categorize_retention_days)
    
    # 计算各分类的频数和占比
    retention_distribution = df['留存天数分类'].value_counts().reset_index()
    retention_distribution.columns = ['留存天数分类', '用户数']
    retention_distribution['占比'] = retention_distribution['用户数'] / retention_distribution['用户数'].sum()
    
    # 按照指定顺序重新排序
    order_mapping = {
        '当天': 1,
        '次日': 2,
        '2-7天': 3,
        '8-30天': 4,
        '30天以上': 5
    }
    retention_distribution['排序'] = retention_distribution['留存天数分类'].map(order_mapping)
    retention_distribution = retention_distribution.sort_values('排序').drop('排序', axis=1)
    
    print("留存天数分布:")
    for _, row in retention_distribution.iterrows():
        print(f"{row['留存天数分类']}: {row['用户数']}人 ({row['占比']*100:.2f}%)")
    
    # ========== 保存到Excel并创建图表 ==========
    print("正在保存到Excel并创建图表...")
    
    with pd.ExcelWriter('付费用户留存分析.xlsx', engine='openpyxl') as writer:
        # 表单1：按自然周留存率分析
        weekly_retention_df.to_excel(writer, sheet_name='按自然周留存率分析', index=False)
        
        # 表单2：留存天数分布
        retention_distribution.to_excel(writer, sheet_name='留存天数分布', index=False)
        
        # 获取工作簿对象
        workbook = writer.book
        
        # ========== 为表单1创建留存率趋势图 ==========
        worksheet1 = workbook['按自然周留存率分析']
        
        # 创建留存率趋势图
        chart1 = LineChart()
        chart1.title = "付费用户留存率趋势（按自然周）"
        chart1.x_axis.title = "自然周"
        chart1.y_axis.title = "留存率"
        chart1.y_axis.scaling.min = 0
        chart1.y_axis.scaling.max = 1
        chart1.y_axis.number_format = '0%'
        
        # 准备数据
        data = Reference(worksheet1, min_col=6, min_row=1, max_row=len(weekly_retention_df)+1, max_col=8)
        cats = Reference(worksheet1, min_col=1, min_row=2, max_row=len(weekly_retention_df)+1)
        
        chart1.add_data(data, titles_from_data=True)
        chart1.set_categories(cats)
        
        # 设置图表大小和位置
        chart1.width = 15
        chart1.height = 10
        worksheet1.add_chart(chart1, "J2")
        
        # ========== 为表单2创建留存天数分布柱状图 ==========
        worksheet2 = workbook['留存天数分布']
        
        # 创建留存天数分布柱状图
        chart2 = BarChart()
        chart2.title = "付费用户留存天数分布"
        chart2.x_axis.title = "留存天数分类"
        chart2.y_axis.title = "用户数量"
        
        # 准备数据
        data2 = Reference(worksheet2, min_col=2, min_row=1, max_row=len(retention_distribution)+1, max_col=2)
        cats2 = Reference(worksheet2, min_col=1, min_row=2, max_row=len(retention_distribution)+1)
        
        chart2.add_data(data2, titles_from_data=True)
        chart2.set_categories(cats2)
        
        # 设置图表大小和位置
        chart2.width = 12
        chart2.height = 8
        worksheet2.add_chart(chart2, "D2")
        
        # 在留存天数分布表格中添加占比列
        # 在D列添加占比数据
        for i, row in enumerate(retention_distribution.iterrows(), start=2):
            percentage = row[1]['占比'] * 100
            worksheet2[f'D{i}'] = f"{percentage:.1f}%"
        
        # 设置D1标题
        worksheet2['D1'] = '占比'
        
        # 调整列宽
        for sheet in workbook.sheetnames:
            ws = workbook[sheet]
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
    
    print(f"分析完成！结果已保存到 '付费用户留存分析.xlsx'")
    
    # 显示详细统计信息
    print(f"\n=== 详细统计信息 ===")
    print(f"平均留存天数: {df['留存天数'].mean():.2f}天")
    print(f"中位数留存天数: {df['留存天数'].median():.2f}天")
    print(f"最长留存天数: {df['留存天数'].max()}天")
    print(f"最短留存天数: {df['留存天数'].min()}天")
    print(f"分析的自然周数: {len(weekly_retention_df)}")
    
    return weekly_retention_df, retention_distribution

if __name__ == "__main__":
    file_path = "3月以来的付费用户情况.xlsx"
    weekly_retention_df, retention_distribution = analyze_retention(file_path) 