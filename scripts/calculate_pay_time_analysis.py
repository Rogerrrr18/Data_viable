import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference

# 读取原始数据
file_path = 'Result_10.xlsx'
raw_df = pd.read_excel(file_path)

# 只保留有首次付费日期的用户
paid_df = raw_df.dropna(subset=['首次付费时间']).copy()
paid_df['注册时间'] = pd.to_datetime(paid_df['注册时间'])
paid_df['首次付费时间'] = pd.to_datetime(paid_df['首次付费时间'])
paid_df['注册日期'] = paid_df['注册时间'].dt.date
# 付费时长（天）
paid_df['付费时长'] = (paid_df['首次付费时间'].dt.date - paid_df['注册日期']).apply(lambda x: x.days)

# 1. 频数分布直方图（天）
bins = list(range(0, paid_df['付费时长'].max()+2))  # 0,1,2,...max+1
bin_labels = [f"{i}-{i+1}天" for i in range(0, paid_df['付费时长'].max()+1)]
hist_data = pd.cut(paid_df['付费时长'], bins=bins, labels=bin_labels, right=False, include_lowest=True)
hist_counts = hist_data.value_counts().sort_index()
hist_df = pd.DataFrame({'付费时长区间': bin_labels, '频数': hist_counts.values})

# 2. 24小时内付费时长频数分布（小时）
paid_df['付费时长_小时'] = (paid_df['首次付费时间'] - paid_df['注册时间']).dt.total_seconds() / 3600
within_24h = paid_df[(paid_df['付费时长_小时'] >= 0) & (paid_df['付费时长_小时'] < 24)]
hour_bins = list(range(0, 25))  # 0,1,...,24
hour_labels = [f"{i}-{i+1}h" for i in range(24)]
hist_hour_data = pd.cut(within_24h['付费时长_小时'], bins=hour_bins, labels=hour_labels, right=False, include_lowest=True)
hist_hour_counts = hist_hour_data.value_counts().sort_index()
hist_hour_df = pd.DataFrame({'付费时长区间(小时)': hour_labels, '频数': hist_hour_counts.values})

# 写入Excel
output_file = 'pay_time_analysis.xlsx'
with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    hist_df.to_excel(writer, sheet_name='付费时长分布', index=False)
    hist_hour_df.to_excel(writer, sheet_name='24小时内付费时长分布(小时)', index=False)

# 1. 付费时长分布（天）柱状图
workbook = load_workbook(output_file)
ws = workbook['付费时长分布']
chart = BarChart()
chart.title = "付费用户注册到首次付费时长频数分布"
chart.x_axis.title = "付费时长区间"
chart.y_axis.title = "频数"
data = Reference(ws, min_col=2, min_row=1, max_row=len(hist_df)+1)
cats = Reference(ws, min_col=1, min_row=2, max_row=len(hist_df)+1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.height = 12
chart.width = 18
ws.add_chart(chart, "E2")

# 2. 24小时内付费时长分布（小时）柱状图
ws_hour = workbook['24小时内付费时长分布(小时)']
chart_hour = BarChart()
chart_hour.title = "24小时内付费用户注册到首次付费时长频数分布(小时)"
chart_hour.x_axis.title = "付费时长区间(小时)"
chart_hour.y_axis.title = "频数"
data_hour = Reference(ws_hour, min_col=2, min_row=1, max_row=len(hist_hour_df)+1)
cats_hour = Reference(ws_hour, min_col=1, min_row=2, max_row=len(hist_hour_df)+1)
chart_hour.add_data(data_hour, titles_from_data=True)
chart_hour.set_categories(cats_hour)
chart_hour.height = 12
chart_hour.width = 18
ws_hour.add_chart(chart_hour, "E2")

workbook.save(output_file)
print(f"分析完成，结果已保存到 {output_file}") 